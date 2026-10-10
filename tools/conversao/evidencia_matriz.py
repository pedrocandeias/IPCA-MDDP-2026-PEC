"""Recover attributed historical excerpts from the citation evidence workbook."""
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path
import re
from urllib.parse import unquote

import fitz
import openpyxl


def doi_key(value):
    value = str(value or '').strip().lower()
    return re.sub(r'^https?://(?:dx\.)?doi\.org/', '', value).rstrip('.,; ')


def read_matrix(path, root, records, norm):
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook['Evidência por citação']
    technical = workbook['Controlo técnico']
    assert sheet.max_row == technical.max_row
    results, report, cache = defaultdict(list), [], {}
    by_title, by_doi = defaultdict(set), defaultdict(set)
    for rid, rec in records.items():
        by_title[norm(rec['title'])].add(rid)
        if rec.get('doi'):
            by_doi[doi_key(rec['doi'])].add(rid)
    for row in range(2, sheet.max_row + 1):
        values = [sheet.cell(row, c).value for c in range(1, 8)]
        tech = [technical.cell(row, c).value for c in range(1, 16)]
        assert values[0] == tech[0] and values[6] == tech[11], row
        title, reference = str(values[2] or ''), str(tech[5] or '')
        dois = {doi_key(v) for v in re.findall(r'10\.\d{4,9}/[^\s<>]+', reference)}
        doi_ids = set().union(*(by_doi[d] for d in dois)) if dois else set()
        clean_title = re.sub(r'\s*(?:\[(?:Conference paper|Computer software)\]|\((?:ISO Standard No\.|IEC Standard No\.|Royal College of Art Research Papers|Report No)).*$', '', title)
        title_ids = by_title[norm(clean_title)]
        ids, method = doi_ids or title_ids, 'DOI' if doi_ids else 'título exato'
        # Reject conflicts rather than quietly attributing another work's excerpt.
        if doi_ids and title_ids and doi_ids != title_ids:
            ids = set()
            method = 'conflito entre DOI e título'
        if not ids and not doi_ids and not title_ids:
            candidates = set()
            for rid, rec in records.items():
                ratio = SequenceMatcher(None, norm(title), norm(rec['title']), autojunk=False).ratio()
                if ratio >= .96 and str(rec.get('year')) == str(values[4]):
                    candidates.add(rid)
            ids, method = candidates, 'título semelhante e ano'
        audit = {'row': row, 'title': title, 'record_id': '', 'match_method': method, 'status': ''}
        report.append(audit)
        if len(ids) != 1:
            audit['status'] = 'obra não resolvida inequivocamente'
            continue
        rid = next(iter(ids))
        audit['record_id'] = rid
        raw = str(values[5] or '')
        match = re.match(r'^\[PDF, p\.\s*(\d+)\]\s*(.+)', raw, re.S)
        if not match:
            audit['status'] = 'sem excerto paginado'
            continue
        page, excerpt = int(match[1]), match[2].strip()
        if page != tech[8]:
            raise ValueError(f'Página divergente na linha {row}')
        link = sheet.cell(row, 3).hyperlink
        pdf = (path.parent / unquote(link.target)).resolve() if link else root/'material/bibliografia'/str(tech[7])
        if not pdf.is_relative_to(root.resolve()) or not pdf.is_file():
            audit['status'] = 'PDF de proveniência indisponível'
            continue
        key = (str(pdf), page)
        if key not in cache:
            with fitz.open(pdf) as document:
                cache[key] = document[page-1].get_text() if 0 < page <= len(document) else ''
        compact = lambda text: norm(text).replace(' ', '')
        verified = compact(excerpt) in compact(cache[key])
        # A recovered audit quotation stays historical when PDF extraction differs.
        # It is never presented as a new assessment of the current manuscript.
        audit['status'] = 'excerto confirmado na página PDF' if verified else 'transcrição histórica; extração PDF não coincidente'
        assessment = str(values[6] or '')
        if tech[14]:
            assessment += '. ' + str(tech[14])
        results[rid].append({
            'section': f'Evidência por citação, linha {row}',
            'claims': [str(values[1] or '')], 'excerpts': [excerpt],
            'pages': f'PDF da fonte, página {page}. DOCX auditado, página {values[0]}.',
            'assessment': assessment, 'assessment_origin': str(tech[12] or ''),
            'pdf': str(pdf.relative_to(root.resolve())),
            'origin': str(path.relative_to(root)), 'origin_kind': 'xlsx',
            'match_method': method, 'audit_reference': reference,
            'excerpt_method': str(tech[10] or ''), 'pdf_text_verified': verified,
        })
    workbook.close()
    return results, report
