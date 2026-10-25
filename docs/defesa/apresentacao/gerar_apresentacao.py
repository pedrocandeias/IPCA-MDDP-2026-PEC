#!/usr/bin/env python3
"""Gera PPTX, gráfico e manifesto a partir das fontes locais da defesa."""
from pathlib import Path
import csv, json, hashlib, sys, shutil
from collections import defaultdict
try:
    import pptx
except ImportError:
    sys.path.insert(0, '/tmp/defesa_pptx_deps')
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from exportar_guiao_falado import ler_falas

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[2]
ASSETS=HERE/'assets'
CSV=ROOT/'suplementos/04_preparacao_impressao_prototipos/registos/comparacao_dimensional_72_medicoes.csv'
ROWS=list(csv.DictReader(CSV.open()))
assert len(ROWS)==72 and all(r['numero_leituras']=='1' for r in ROWS)
PHOTO=ROOT/'suplementos/04_preparacao_impressao_prototipos/fotografias_originais/teste-impressao-phoenix-15-anos-13221117.jpg'
shutil.copy2(PHOTO,ASSETS/'phoenix_montagem.jpg')
BG='F6F5F0'; INK='193536'; TEAL='007C78'; MUTED='566B6A'; LIGHT='E3ECE8'; WHITE='FFFFFF'; ORANGE='AD542D'
FONT='Liberation Sans'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='HandFab: Defesa do projeto de mestrado'
prs.core_properties.subject='Design e Desenvolvimento de Produto · UPCA'
prs.core_properties.author='Pedro Miguel Candeias da Silva'
records=[]
MAIN_SLIDES=23
visuals=[]
t20=[35, 75, 60, 40, 65, 80, 55, 90, 40, 45, 45, 45, 60, 55, 45, 20, 30, 50, 40, 75, 45, 85, 20]
t30=[40, 110, 70, 60, 110, 125, 85, 150, 50, 55, 65, 65, 85, 85, 70, 30, 45, 65, 60, 120, 70, 155, 30]
assert sum(t20)==1200 and sum(t30)==1800

def rgb(c):return RGBColor.from_string(c)
def box(s,x,y,w,h,fill=LIGHT,line=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=rgb(fill)
    if line:sh.line.color.rgb=rgb(line)
    else:sh.line.fill.background()
    sh._element.spPr.append(OxmlElement("a:effectLst"))
    return sh

def text(s,x,y,w,h,value,size=24,color=INK,bold=False,font=FONT,align=None):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame
    tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(0);tf.margin_top=tf.margin_bottom=Inches(.025)
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
        p.space_after=Pt(7);p.line_spacing=1.1
        if align is not None:p.alignment=align
    return sh

def picture(s,path,x,y,w,h,bg=WHITE):
    box(s,x,y,w,h,bg)
    with Image.open(path) as im:iw,ih=im.size
    scale=min(w/iw,h/ih);pw,ph=iw*scale,ih*scale
    s.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))

def slide(title,kicker,source,dark=False):
    i=len(records)+1;s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(INK if dark else BG)
    fg=WHITE if dark else INK
    box(s,.45,.45,.12,.3,TEAL)
    text(s,.7,.4,11.8,.3,kicker.upper(),11,TEAL if not dark else '8DCCC1',True)
    text(s,.7,.94,11.9,1.0,title,32,fg,True)
    text(s,.7,6.96,11.55,.26,source,10,'B5CBC6' if dark else MUTED)
    text(s,12.25,6.93,.4,.3,str(i).zfill(2),11,fg,True,align=PP_ALIGN.RIGHT)
    records.append({'numero':i,'titulo':title,'principal':i<=MAIN_SLIDES,'tempo_20_segundos':t20[i-1] if i<=MAIN_SLIDES else None,'tempo_30_segundos':t30[i-1] if i<=MAIN_SLIDES else None,'fonte':source})
    return s

def card(s,x,y,w,num,title,body,h=3.4):
    box(s,x,y,w,h,WHITE)
    text(s,x+.22,y+.18,w-.44,.48,num,28,TEAL,True)
    text(s,x+.22,y+.87,w-.44,.85,title,23,INK,True)
    text(s,x+.22,y+1.8,w-.44,h-1.95,body,19,MUTED)

def row(s,y,a,b,c=None):
    box(s,.7,y,11.9,.79,WHITE)
    text(s,.95,y+.1,2.65,.58,a,20,TEAL,True)
    text(s,3.8,y+.1,4.15 if c else 8.25,.6,b,18)
    if c:text(s,8.2,y+.1,4.1,.6,c,18,MUTED)


def region(s,name,x,y,w,h,crop=(0,0,1,1)):
    """Enquadra uma região sem reamostrar a imagem incorporada nem a deformar."""
    path=ASSETS/name
    with Image.open(path) as im: iw,ih=im.size
    left,top,right,bottom=crop
    cw,ch=iw*(right-left),ih*(bottom-top)
    scale=min(w/cw,h/ch);pw,ph=cw*scale,ch*scale
    sh=s.shapes.add_picture(str(path),Inches(x+(w-pw)/2),Inches(y+(h-ph)/2),width=Inches(pw),height=Inches(ph))
    sh.crop_left=left;sh.crop_top=top;sh.crop_right=1-right;sh.crop_bottom=1-bottom
    visuals.append({'diapositivo':len(records),'ficheiro':'assets/'+name,'regiao_normalizada':[left,top,right,bottom]})
    return sh

def screen(s,name,crop=(.10,.10,.94,.98)):
    region(s,name,.7,1.95,11.9,4.53,crop)


# Chart: ranges across all three axes and four profiles, no inference beyond observed values.
fig,ax=plt.subplots(figsize=(10.7,4.5));fig.patch.set_facecolor('#'+BG);ax.set_facecolor('#'+BG)
groups=[]
for mi,m in enumerate(['Flexy Beast','Paraglider Hand','UnLimbited Phoenix']):
 for mat,color,offset in [('PLA','#'+TEAL,.13),('PETG','#'+ORANGE,-.13)]:
  vals=[float(r['desvio_malha_peca_mm']) for r in ROWS if r['modelo']==m and r['material']==mat]
  assert len(vals)==12
  y=2-mi+offset;lo,hi=min(vals),max(vals)
  ax.plot([lo,hi],[y,y],color=color,lw=5,solid_capstyle='round',label=mat if mi==0 else None)
  ax.scatter([lo,hi],[y,y],s=42,color=color,zorder=3)
  groups.append({'modelo':m,'material':mat,'n':len(vals),'min_mm':lo,'max_mm':hi})
ax.set_yticks([2,1,0],['Flexy Beast','Paraglider Hand','UnLimbited Phoenix'],fontsize=16)
ax.set_xlim(-.6,.025);ax.set_ylim(-.55,2.6);ax.axvline(0,color='#'+MUTED,lw=1)
ax.set_xticks([-.6,-.4,-.2,0],['−0,6','−0,4','−0,2','0']);ax.tick_params(axis='x',labelsize=15);ax.tick_params(axis='y',length=0,pad=12)
ax.set_xlabel('Dimensão física − dimensão da malha (mm)',fontsize=16,labelpad=12)
ax.spines[['top','right','left']].set_visible(False);ax.spines['bottom'].set_color('#'+MUTED)
ax.grid(axis='x',alpha=.2);ax.legend(loc='upper left',bbox_to_anchor=(0,1.18),frameon=False,ncol=2,fontsize=15)
fig.subplots_adjust(left=.29,right=.97,top=.81,bottom=.2)
fig.savefig(ASSETS/'desvios_dimensionais.png',dpi=190);fig.savefig(ASSETS/'desvios_dimensionais.svg');plt.close(fig)
(HERE/'resumo_dimensional.json').write_text(json.dumps(groups,ensure_ascii=False,indent=2)+'\n')

s=slide('Da medida à peça física', 'Mestrado em Design e Desenvolvimento de Produto · UPCA', 'Imagem: produção do autor · Suplemento 4 · Phoenix, observação de montagem em bancada.')
text(s,.7,2.04,6.15,2.7,'Integração do design e da inteligência artificial em processos paramétricos para o desenvolvimento de próteses de membros superiores em impressão 3D',23,INK,True)
text(s,.7,4.95,6.1,.65,'Pedro Miguel Candeias da Silva',21,INK,True)
text(s,.7,5.63,6.1,.9,'Orientador: Doutor Demétrio Ferreira Matos\nProjeto de mestrado · 2026',17,MUTED)
s.shapes.add_picture(str(ASSETS/'upca_logo_oficial.png'), Inches(9.7), Inches(.45), width=Inches(2.9))
region(s,'phoenix_montagem.jpg',7.1,2.02,5.5,3.76)
text(s,7.12,5.94,5.4,.52,'HandFab: da configuração à peça',18,TEAL,True)

s=slide('O que motivou o desenvolvimento do projeto?', 'Motivação · Acesso e personalização', 'Manuscrito, 1.1–1.2 e 2.8, Lacuna 5 · Andrysek (2010); Alluhydan et al. (2023); Baumann e Maria (2023).')
text(s,.85,1.97,11.5,.9,'Dificuldades de acesso a próteses adaptadas,\nparticularmente em contextos de baixos e médios rendimentos.',26,INK,True)
box(s,.7,3.12,5.8,2.17,WHITE)
text(s,.95,3.32,5.3,.4,'Barreiras identificadas',23,TEAL,True)
text(s,.95,3.93,5.3,1.15,'Custo e apoio técnico especializado\nAdaptação, manutenção e infraestrutura',21)
box(s,6.8,3.12,5.8,2.17,WHITE)
text(s,7.05,3.32,5.3,.4,'Possibilidade a explorar',23,TEAL,True)
text(s,7.05,3.93,5.3,1.15,'Modelos abertos e impressão 3D\nConfiguração ajustável e compreensível',21)
box(s,.7,5.57,11.9,1.06,LIGHT)
text(s,.95,5.64,11.35,.9,'Desenvolver o HandFab para ligar medidas a modelos configuráveis\ne examinar o resultado até às peças físicas.',22,INK,True)

s=slide('Como passar de uma medida a uma peça?', 'Problema de design', 'Captura do HandFab e fotografias do autor · Configuração e peças ilustram etapas do projeto.')
region(s,'flexy_parametros.png',.7,2.08,7.1,3.87)
region(s,'flexy_perfis.jpeg',8.02,2.08,4.58,3.87)
text(s,.8,6.16,7.0,.45,'Que efeito tem o valor que introduzo?',24,TEAL,True)
text(s,8.1,6.16,4.5,.45,'O que chega à peça?',24,TEAL,True)

s=slide('Perguntas de investigação e âmbito da avaliação', 'Perguntas e âmbito', 'Síntese das perguntas aprovadas · Formulações integrais no diapositivo 24 · Tabela 31.')
card(s,.7,2.05,3.8,'P1','Configuração e controlo','Como apoiar escolhas sobre o modelo?')
card(s,4.75,2.05,3.8,'P2','Avaliação e reprodução','Como examinar o processo e os resultados?')
card(s,8.8,2.05,3.8,'P3','Requisitos e experiência','Como responder às necessidades das pessoas?')
text(s,.9,5.95,11.5,.52,'Âmbito realizado: verificação técnica e dimensional, sem participantes.',22,TEAL,True)
# Objetivos explícitos, sintetizados sem eliminar os sete pontos do manuscrito.

s=slide('Objetivo geral e objetivos específicos', 'Objetivos da investigação', 'Síntese da secção 1.3 do manuscrito · Sete objetivos específicos, com o âmbito original preservado.')
box(s,.7,1.92,11.9,1.62,LIGHT)
text(s,.92,2.04,11.4,.26,'OBJETIVO GERAL',12,TEAL,True)
text(s,.92,2.35,11.4,1.12,'Produzir conhecimento através da prática de design, desenvolvendo e avaliando tecnicamente um sistema paramétrico apoiado por IA para a configuração preliminar de modelos de mão protésica destinados ao fabrico aditivo.',20,INK,True)
text(s,.92,3.64,11.4,.28,'OBJETIVOS ESPECÍFICOS',12,TEAL,True)
objectives=[
 'Analisar como o design traduz requisitos humanos em decisões de projeto e fabrico.',
 'Organizar dados antropométricos, preservando a sua caracterização e limitações.',
 'Adaptar modelos abertos de mão a uma estrutura paramétrica explícita em OpenSCAD.',
 'Integrar o percurso na web, com visualização e gestão de versões, até à impressão.',
 'Delimitar a IA à sugestão inicial, distinguindo regras geométricas e decisões humanas.',
 'Avaliar a coerência entre dados, parâmetros, geometria digital e protótipos físicos.',
 'Definir condições para avaliação clínica posterior, incluindo função e utilização.',
]
for i,objective in enumerate(objectives):
 x=.92 if i<4 else 7.02; y=4.12+(i if i<4 else i-4)*.66
 text(s,x,y,.42,.34,f'{i+1:02d}',16,TEAL,True)
 text(s,x+.48,y,5.12,.64,objective,17)

# Metodologia e ferramentas de investigação no percurso principal.

s=slide('Como organizei a investigação e o desenvolvimento', 'Metodologia de investigação', 'Manuscrito, 3.1–3.3 e Tabela 4 · Frayling (1994); Zimmerman et al. (2007); Design Council (2007).')
text(s,.9,1.98,11.5,.45,'Double Diamond: organizar a exploração e a convergência',24,TEAL,True)
phases=[('01','Descobrir','Literatura e modelos\nNecessidades'),
        ('02','Definir','Problema e requisitos\nÂmbito da avaliação'),
        ('03','Desenvolver','Modelos paramétricos\nPlataforma web'),
        ('04','Entregar','Protótipos\nVerificação técnica')]
for i,(num,title,body) in enumerate(phases):
 x=.7+i*3.05
 box(s,x,2.66,2.75,1.98,WHITE)
 text(s,x+.17,2.82,2.4,.32,num,15,TEAL,True)
 text(s,x+.17,3.23,2.4,.4,title,23,INK,True)
 text(s,x+.17,3.78,2.4,.8,body,17,MUTED)
 if i<3: text(s,x+2.77,3.35,.27,.4,'›',24,TEAL,True)
text(s,.9,4.84,11.5,.4,'Os ensaios podem levar à revisão de decisões anteriores.',20,MUTED)
box(s,.7,5.44,11.9,1.2,LIGHT)
text(s,.93,5.57,11.4,.4,'Research Through Design: produzir conhecimento através do projeto',22,TEAL,True)
text(s,.93,6.12,11.4,.33,'Problema observado → decisão → protótipo → ensaio → reflexão e alteração',20,INK)

s=slide('Como integrei a IA na investigação', 'Ferramentas de apoio e objeto de estudo', 'Declaração de utilização de IA e capítulo 6 · Registos locais da pesquisa bibliográfica e do desenvolvimento.')
box(s,.7,2.02,5.8,3.64,WHITE)
text(s,.95,2.22,5.3,.45,'No trabalho de investigação',24,TEAL,True)
text(s,.95,2.94,5.3,.85,'Elicit\nPesquisa e agregação de bibliografia',21,INK,True)
text(s,.95,4.14,5.3,1.15,'ChatGPT, Codex e Claude\nApoio à organização, redação, revisão\ne desenvolvimento do código',20,MUTED)
box(s,6.8,2.02,5.8,3.64,WHITE)
text(s,7.05,2.22,5.3,.45,'Dentro do HandFab',24,TEAL,True)
text(s,7.05,2.94,5.3,.85,'Sugestões iniciais\nValores para parâmetros existentes',21,INK,True)
text(s,7.05,4.14,5.3,1.15,'Revisão das propostas\nAplicação de regras geométricas\nVerificação dos efeitos na peça',20,MUTED)
text(s,.94,5.85,11.4,.85,'O investigador seleciona as fontes, verifica os resultados\ne assume a responsabilidade pelas decisões.',22,TEAL,True)

s=slide('Elicit: da pergunta à bibliografia organizada', 'Pesquisa e revisão da literatura', 'Registos Elicit: pesquisas, relatórios e sessões por secção · Exemplo documentado no pacote da secção 6.1.')
steps=[('01','Perguntar','Delimitar o tema\ne a secção'),
       ('02','Pesquisar','Localizar estudos\ncom o Elicit'),
       ('03','Agregar','Referências e\nsínteses temáticas'),
       ('04','Verificar','Consultar originais\nAvaliar pertinência'),
       ('05','Integrar','Organizar fontes\nRedigir a síntese')]
for i,(num,title,body) in enumerate(steps):
 x=.7+i*2.43
 box(s,x,2.11,2.18,2.13,LIGHT if i==3 else WHITE)
 text(s,x+.16,2.27,1.86,.34,num,17,TEAL,True)
 text(s,x+.16,2.78,1.86,.4,title,21,INK,True)
 text(s,x+.16,3.37,1.86,.73,body,16,MUTED)
 if i<4: text(s,x+2.18,2.86,.25,.4,'›',23,TEAL,True)
text(s,.94,4.53,11.35,.34,'EXEMPLO: PAPEL DA IA NA CONFIGURAÇÃO PROTÉSICA',13,TEAL,True)
text(s,.94,5.04,11.35,.76,'Distinguir estudos sobre controlo da prótese de estudos\nsobre apoio à configuração a partir de dados antropométricos.',24,INK,True)
text(s,.94,6.16,11.35,.4,'Registo local: pergunta, resultados, relatório, referências e notas por secção.',19,MUTED)

# Referências que organizam o argumento

s=slide('Como a literatura orienta o projeto', 'Fundamentação e seleção', 'Referências já mobilizadas no manuscrito · Seleção para esta exposição pela função no argumento.')
items=[('Cross (1982)','Conhecimento de design','Aprendizagem a partir das decisões de projeto'),
       ('Frayling (1994); Zimmerman et al. (2007)','Investigação através do projeto','Análise documentada do artefacto'),
       ('Fischer et al. (2004)','Meta-design e modificação','Escolhas disponíveis no HandFab'),
       ('Manero et al. (2019); Cabibihan et al. (2021)','Próteses impressas em 3D','Contexto e necessidade de avaliação')]
for i,(a,b,c) in enumerate(items):
 y=2.12+i*.94;box(s,.7,y,11.9,.84,WHITE)
 text(s,.91,y+.12,3.8,.6,a,19,TEAL,True)
 text(s,4.85,y+.1,3.4,.62,b,18,INK,True)
 text(s,8.4,y+.1,3.95,.62,c,18,MUTED)
text(s,.94,6.2,11.4,.5,'Critério: utilidade para justificar decisões e interpretar os resultados.',21,INK,True)


s=slide('HandFab: um percurso com decisões explícitas', 'Escolher um modelo', 'HandFab · Catálogo documentado na Figura 22 do manuscrito · Captura do projeto.')
screen(s,'figura_22_1.png',(.11,.11,.90,.77))
text(s,.88,6.46,11.8,.4,'Quatro modelos disponíveis para iniciar a configuração.',22,TEAL,True)

s=slide('Fontes antropométricas e critérios de utilização', 'Organizar os dados de referência', 'Gordon et al. (1989; 2014) · TU Delft, DINED (s.d.) · ISO 7250-1:2017 · Figura 24.')
region(s,'figura_24_1.png',.7,1.97,7.65,4.77,(.13,.035,.90,.96))
text(s,8.75,2.23,3.6,.8,'Guardar a origem de cada valor',26,TEAL,True)
text(s,8.75,3.5,3.55,2.5,'População e idade\nMedida e unidade\nEstatística utilizada\nFonte e protocolo',23)

s=slide('Da base consolidada aos perfis do HandFab', 'Dados na plataforma', 'Manuscrito, 4.2 e Figura 23 · Suplemento 1 · Contagens de perfis da aplicação.')
region(s,'figura_23_1.png',.7,1.96,8.85,4.82,(.035,.105,.985,.96))
text(s,9.96,2.25,2.64,.85,'100 perfis',32,TEAL,True)
text(s,9.96,3.45,2.64,2.3,'97 em lote\n2 por CSV\n1 manual',23)
text(s,9.96,5.82,2.64,.7,'Referências estatísticas',18,MUTED)

s=slide('Cada modelo exige regras próprias', 'Decisões de projeto', 'Manuscrito, 4.4 e 8.3 · Fotografias ilustrativas dos modelos; não constituem comparação antes/depois.')
region(s,'paraglider_perfis.jpeg',.7,2.0,5.8,3.3)
region(s,'phoenix_montagem.jpg',6.8,2.0,5.8,3.3)
text(s,.87,5.48,5.45,.45,'Paraglider Hand',25,TEAL,True)
text(s,.87,6.04,5.45,.64,'Corrigir a aplicação da escala na palma.',20)
text(s,6.97,5.48,5.45,.45,'UnLimbited Phoenix',25,TEAL,True)
text(s,6.97,6.04,5.45,.64,'Largura mínima de cerca de 82 mm no caso estudado.',20)

s=slide('A IA propõe um ponto de partida', 'Introduzir um pedido', 'HandFab · Figura 22 · Captura existente da interface de assistência, sem nova execução de IA.')
screen(s,'figura_22_2.png',(.10,.14,.95,.77))
text(s,.88,6.46,11.8,.4,'Rever os valores sugeridos antes de gerar a geometria.',22,TEAL,True)

s=slide('Configuração e pré-visualização no HandFab', 'Ajustar os parâmetros', 'HandFab · Figura 22 · Controlos de configuração e pré-visualização da Phoenix.')
screen(s,'figura_22_3.png',(.10,.145,.95,.99))
text(s,.88,6.46,11.8,.4,'Relacionar cada controlo com a alteração da forma.',22,TEAL,True)

s=slide('Escolher cores e materiais no HandFab', 'Configuração do modelo', 'HandFab · Figura 22 · Opções de cor e material apresentadas na interface.')
screen(s,'figura_22_4.png',(.10,.145,.95,.99))
text(s,.88,6.46,11.8,.4,'Ver as escolhas de cor em cada componente.',22,TEAL,True)

s=slide('Da geometria gerada à exportação', 'Preparar a passagem para o fabrico', 'HandFab · Figura 20 · Código OpenSCAD, pré-visualização e comando de exportação.')
region(s,'figura_20_1.png',.7,1.96,8.2,4.8)
text(s,9.2,2.4,3.3,1.2,'Examinar a forma gerada',26,TEAL,True)
text(s,9.2,4.34,3.3,1.3,'Exportar a malha para preparar a impressão',23)

s=slide('A avaliação acompanhou etapas diferentes', 'Impressão e observação das peças', 'Suplemento 4 · Fotografias originais na placa da Bambu Lab A1 e da palma impressa.')
region(s,'flexy_placa.jpg',.7,1.98,3.62,4.33)
region(s,'falanges_placa.jpg',4.49,1.98,3.62,4.33)
region(s,'palma_impressa.jpg',8.28,1.98,4.32,4.33)
text(s,.86,6.46,11.74,.36,'4 projetos na Série A · 12 casos na Série B · 72 medições físicas',21,TEAL,True)

s=slide('Comparar peças de diferentes perfis', 'Séries físicas de ensaio', 'Fotografias do autor · Suplemento 4 · Idades dos perfis de referência, sem participantes.')
region(s,'flexy_perfis.jpeg',.7,2.02,5.8,3.63)
region(s,'paraglider_perfis.jpeg',6.8,2.02,5.8,3.63)
text(s,.87,5.8,5.45,.4,'Flexy Beast',25,TEAL,True)
text(s,6.97,5.8,5.45,.4,'Paraglider Hand',25,TEAL,True)
text(s,.88,6.43,11.5,.4,'Perfis de referência: 8, 15, 28 e 70 anos.',22)

s=slide('Peça e malha: diferenças medidas, com limites', 'Resultados dimensionais', 'Fonte: Suplemento 4, comparacao_dimensional_72_medicoes.csv · Gráfico calculado a partir dos registos.')
picture(s,ASSETS/'desvios_dimensionais.png',.7,1.95,11.9,4.26,bg=BG)
text(s,.92,6.25,11.5,.47,'Amplitude por grupo, calculada com uma leitura por combinação.',21,TEAL,True)

s=slide('Phoenix: uma montagem observada em bancada', 'Montagem do protótipo', 'Suplemento 4 · PHX-TEEN15-MONT-001 · Um espécime PLA, perfil de 15 anos; observação em bancada.')
region(s,'phoenix_montagem.jpg',.7,1.97,8.8,4.55)
region(s,'phoenix_articulacao.jpeg',9.72,1.97,2.88,4.55)
text(s,.87,6.5,11.5,.36,'Montagem e movimento manual · Pormenor da articulação',21,TEAL,True)

s=slide('O contributo para a configuração do produto', 'Resposta às perguntas', 'Síntese da Tabela 31 e de 9.1-9.4 · Cross (1982): enquadramento do conhecimento de design.')
text(s,.94,2.0,2.55,.35,'EIXO',12,TEAL,True);text(s,3.8,2.0,4.0,.35,'EVIDÊNCIA OBTIDA',12,TEAL,True);text(s,8.2,2.0,4.1,.35,'POR AVALIAR',12,ORANGE,True)
row(s,2.65,'Configuração','Valores editáveis e revisão','Conforto e adequação individual')
row(s,3.7,'Avaliação','Verificação técnica e dimensional','Usabilidade e durabilidade')
row(s,4.75,'Requisitos','Regras e escolhas disponíveis','Aceitação e autonomia')
text(s,.94,6.0,11.2,.56,'Explicar como cada escolha de configuração afeta a forma.',26,INK,True)

s=slide('As regras de cada modelo condicionam o resultado', 'Conclusão', 'Projeto HandFab · Pedro Miguel Candeias da Silva · Obrigado.', dark=True)
region(s,'phoenix_montagem.jpg',.7,1.96,7.6,4.5)
text(s,8.7,2.5,3.84,.5,'Próximo passo',26,'8DCCC1',True)
text(s,8.7,3.43,3.84,2.4,'Avaliar o HandFab com as pessoas que o irão utilizar.',28,WHITE,True)

s=slide('Perguntas de investigação aprovadas', 'Apoio · Formulações integrais', 'Fonte: manuscrito, 9.1. Diapositivo fora do percurso principal de 20 minutos.')
qs=['Como o design de produto, métodos paramétricos e inteligência artificial podem melhorar a personalização, o conforto e a adequação funcional de próteses, mantendo a acessibilidade e o controlo sobre as decisões de projeto?', 'Quais metodologias e ferramentas validam a eficácia, a usabilidade, a durabilidade e a possibilidade de reproduzir os resultados de próteses impressas em 3D?', 'Como é que o design industrial concilia requisitos anatómicos, funcionais, ergonómicos, estéticos e simbólicos, promovendo a aceitação, a dignidade e a autonomia?']
for i,q in enumerate(qs):
 text(s,.78,2.03+i*1.47,.65,.5,f'P{i+1}',23,TEAL,True);text(s,1.62,2.03+i*1.47,10.8,1.3,q,22)

s=slide('Os números representam conjuntos diferentes', 'Apoio · Desenho da avaliação', 'Fonte: capítulo 8, Anexo D e CSV das Séries A/B e medições físicas.')
card(s,.7,2.08,3.8,'4','Projetos da Série A','Condições diferentes. Comparação direta entre equipamentos limitada.')
card(s,4.75,2.08,3.8,'12','Casos da Série B','3 modelos × 4 perfis. Preparação digital concluída, com avisos.')
card(s,8.8,2.08,3.8,'72','Medições físicas','3 modelos × 4 perfis × 3 eixos × 2 materiais. Uma leitura por combinação.')
text(s,.96,6.08,11.2,.45,'Perfis de ensaio: 8, 15, 28 e 70 anos. Não correspondem a participantes.',21,TEAL,True)

s=slide('Valores registados: amplitude por grupo', 'Apoio · Comparação dimensional', 'Fonte: comparacao_dimensional_72_medicoes.csv · Diferença = dimensão física − dimensão da malha.')
text(s,.9,1.93,4,.35,'MODELO / MATERIAL',13,TEAL,True);text(s,6,1.93,3,.35,'MÍNIMO (mm)',13,TEAL,True);text(s,9,1.93,3,.35,'MÁXIMO (mm)',13,TEAL,True)
for i,g in enumerate(groups):
 y=2.5+i*.61;box(s,.7,y,11.9,.54,WHITE)
 text(s,.95,y+.04,5,.42,g['modelo']+' · '+g['material'],18)
 text(s,6,y+.04,2.5,.42,f"{g['min_mm']:.3f}".replace('.',','),19,TEAL,True)
 text(s,9,y+.04,2.5,.42,f"{g['max_mm']:.3f}".replace('.',','),19,TEAL,True)

s=slide('As falhas também definem o próximo passo', 'Apoio · Limites encontrados', 'Fonte: 8.1.2 e Anexo B · Ensaios de recuperação e auditoria de acessibilidade.')
card(s,.7,2.08,5.8,'Controlo','Valores fora do limite','Um valor simulado acima do limite chegou ao estado interno.',3.7)
card(s,6.8,2.08,5.8,'Acessibilidade','Barreiras no percurso','5 de 12 itens falharam na verificação manual.',3.7)
text(s,.95,6.13,11.3,.43,'Problemas identificados orientam revisões concretas do sistema.',23,TEAL,True)

s=slide('Repetição e acessibilidade: resultados delimitados', 'Apoio · Verificações da plataforma', 'Fonte: 8.1.2 e Anexo B. Repetições idênticas dentro da configuração examinada.')
card(s,.7,2.14,5.8,'7 / 5 / 5','Exportações concluídas','Flexy Beast / Paraglider / Phoenix. Objetivo inicial: 10 por modelo.',3.6)
card(s,6.8,2.14,5.8,'7 conformes · 5 falhas','Verificação manual','12 itens. A auditoria identificou barreiras a corrigir.',3.6)

# Referências de apoio: design

s=slide('Referências: design e investigação', 'Apoio · Referências', 'Seleção da bibliografia do manuscrito · A ligação ao HandFab é uma interpretação aplicada ao projeto.')
refs=[('Cross (1982)','Designerly ways of knowing. Design Studies, 3(4), 221-227. DOI: 10.1016/0142-694X(82)90040-0'),
('Fischer et al. (2004)','Meta-design: A manifesto for end-user development. Communications of the ACM, 47(9), 33-37. DOI: 10.1145/1015864.1015884'),
('Frayling (1994)','Research in art and design. Royal College of Art Research Papers, 1(1).'),
('Zimmerman et al. (2007)','Research through design as a method for interaction design research in HCI. CHI, 493-502. DOI: 10.1145/1240624.1240704'),
('Design Council (2007)','Eleven lessons: Managing design in eleven global brands: A study of the design process.')]
for i,(a,b) in enumerate(refs):
 y=2.02+i*.89
 text(s,.8,y,3.05,.7,a,18,TEAL,True);text(s,4.03,y,8.48,.78,b,17)
# Referências de apoio: domínio e dados

s=slide('Referências: próteses e antropometria', 'Apoio · Referências', 'Dados e atribuições: Suplemento 1 · Ensaios do projeto: Suplementos 2-4 · Obras distintas com funções próprias.')
refs=[('Manero et al. (2019)','Implementation of 3D printing technology in the field of prosthetics: Past, present, and future. DOI: 10.3390/ijerph16091641'),
('Cabibihan et al. (2021)','Suitability of the openly accessible 3D printed prosthetic hands for war-wounded children. DOI: 10.3389/frobt.2020.594196'),
('Gordon et al. (1989; 2014)','1988 / 2012 anthropometric survey of U.S. Army personnel: Methods and summary statistics. Relatórios distintos: ANSUR 1988 / ANSUR II.'),
('TU Delft · DINED (s.d.)','DINED: 1D anthropometric database [Base de dados]. Consulta: 02/05/2026. https://dined.io.tudelft.nl/en/database/tool'),
('ISO (2017)','ISO 7250-1:2017. Basic human body measurements for technological design: Part 1: Body measurement definitions and landmarks (2.ª ed.).')]
for i,(a,b) in enumerate(refs):
 y=2.02+i*.89
 text(s,.8,y,3.05,.7,a,18,TEAL,True);text(s,4.03,y,8.48,.78,b,17)

falas=ler_falas()
assert len(falas)==len(records)==len(prs.slides)==30
for s, record in zip(prs.slides, records):
    n=record['numero']; f=falas[n]
    assert record['titulo']==f['titulo'], (n, record['titulo'], f['titulo'])
    record.update(f)
    tempo=(f"{t20[n-1]//60}:{t20[n-1]%60:02d} no percurso principal; "
           f"{t30[n-1]//60}:{t30[n-1]%60:02d} no treino desenvolvido") if n<=MAIN_SLIDES else 'apoio à discussão'
    s.notes_slide.notes_text_frame.text=(
        f'Tempo para me orientar: {tempo}. Confirmar no ensaio.\n\n'
        f'O QUE QUERO DIZER\n{f["fala"]}\n\n'
        + (f'SE PRECISAR DE DESENVOLVER\n{f["aprofundamento"]}\n\n' if f['aprofundamento'] else '')
        + f'LEMBRAR-ME (não ler ao júri)\n{f["indicacao"]}\n\nFONTES PARA CONSULTA\n{record["fonte"]}')

out=HERE/'apresentacao_defesa.pptx';prs.save(out)
(HERE/'diapositivos.json').write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n')
sources=[ASSETS/'upca_logo_oficial.svg',ASSETS/'upca_logo_oficial.png',ROOT/'pedro-candeias-projeto-mestrado-mdddp-ipca-2026-revisto_v1_06102026-0210.odt',ROOT/'a30759_pedro-candeias-projeto-mestrado-mdddp-ipca-v10060211.pdf',ROOT/'pedro-candeias-projeto-mestrado-mdddp-ipca-2026-revisto_v1_23102026-1047.odt',ROOT/'docs/defesa/guiao_apresentacao.md',ROOT/'material/investigacao/elicit/capitulo-6/6.1/searches/search-03.md',ROOT/'material/investigacao/elicit/capitulo-6/6.1/reports/5aaf1fab-74a5-4a0b-a871-eadf2ec0e219/notes.md',ROOT/'material/investigacao/elicit/agent-sessions/INDEX.md',HERE/'guiao_falado.md',ROOT/'suplementos/01_dados_antropometricos/descricao_do_suplemento_1.md',ROOT/'suplementos/01_dados_antropometricos/multi_population_hand.csv',CSV,PHOTO,ROOT/'suplementos/04_preparacao_impressao_prototipos/registos/resultados_serie_b.csv',ROOT/'suplementos/04_preparacao_impressao_prototipos/registos/montagem_phoenix_15_anos.csv']
manifest={'data':'2026-10-25','estado':'Versão 1.6: motivação do desenvolvimento, metodologia e pesquisa bibliográfica com IA e Elicit; sujeita a ensaio','formato':'16:9','diapositivos_principais':MAIN_SLIDES,'apoio':7,'tempo_principal_segundos':sum(t20),'tempo_treino_segundos':sum(t30),'fontes':[{'caminho':str(p.relative_to(ROOT)),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sources],'fotografias':'Capturas existentes do HandFab e fotografias originais do Suplemento 4. Ficheiros integrais incorporados; enquadramentos proporcionais registados em imagens_diapositivos.json.','grafico':'Amplitude min/max de desvio_malha_peca_mm por modelo/material; 12 valores por grupo; sem inferência estatística.','pendencias':['Confirmar condições e data da defesa.','Alinhar as conclusões do manuscrito com o registo de montagem Phoenix.','Conferir a entrada bibliográfica incompleta de Frayling, preservada conforme o manuscrito.','Ensaiar e ajustar o tempo; não houve ensaio cronometrado nem gravação de demonstração.']}
(HERE/'manifesto.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
print(out, len(prs.slides),'diapositivos; notas em todos; percurso principal 20 min; treino 30 min.')

(HERE/'imagens_diapositivos.json').write_text(json.dumps(visuals,ensure_ascii=False,indent=2)+'\n')
